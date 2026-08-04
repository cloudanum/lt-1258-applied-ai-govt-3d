"""Appendix corrections for rev a3 — AppB / AppC / AppD (corrections only, no
enrichment, no slide-count change). All fixes trace to
1258/research/refs-ch05-security-governance.md and refs-ch07-llm-apis-rag-agents.md
(URLs verified live 2026-08-01).

AppB (22 slides, stays 22):
  s2  SAIF six-element list garbled -> official six per Google (ch05 Corrections #6)
  s11 Perspective API 'Do Now' dies Dec 31, 2026 -> rubric exercise (ch05 Corrections #3)
  s13 note: 'CISA ZTMM modeled after BeyondCorp' -> aligned to OMB M-22-09 (#13)
AppC (29 slides, stays 29):
  s14 stale 'future decoders for BERT' -> the encoder/decoder fork (ch07 Corrections #8)
  s26 'continuously learns from new data' -> frozen weights; context/RAG (ch07 #9)
AppD (20 slides, stays 20):
  s10 Word2Vec 'one hidden layer' -> projection layer, no non-linear hidden (#7)
  s11 GloVe TF-IDF claim -> co-occurrence matrix + weighted least squares (#6)
"""
from pathlib import Path
import copy
import pptx_tools as P

DECKS = Path("/Users/iahmad/Creator/Courses_and_conferences/LT/1258/1258a4-author-input/decks")


# --------------------------------------------------------------------------- #
# small in-place text-frame helpers (edit run.text to preserve formatting)
# --------------------------------------------------------------------------- #
def set_par(p, text, level=None):
    """Set a paragraph's text into its first run (preserving run formatting);
    drop any extra runs. Optionally set indent level."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.text = text
    if level is not None:
        p.level = level


def del_par(p):
    p._element.getparent().remove(p._element)


def clone_par_after(p_src, text):
    """Deep-copy p_src (keeps its bullet/level formatting), insert right after
    it, set the copy's text, and return the new paragraph."""
    el = copy.deepcopy(p_src._element)
    p_src._element.addnext(el)
    from pptx.text.text import _Paragraph
    newp = _Paragraph(el, p_src._parent)
    set_par(newp, text)
    return newp


def body_shape(slide, name="Content Placeholder"):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith(name):
            return sh
    raise LookupError(name)


def set_notes(slide, notes):
    slide.notes_slide.notes_text_frame.text = notes


# --------------------------------------------------------------------------- #
def fix_appb():
    path = DECKS / "1258-AppB-Security-DeepDive.pptx"
    prs = P.open_deck(path)
    slides = list(prs.slides)

    # --- slide 2 (idx 1): SAIF official six core elements ------------------ #
    s = slides[1]
    tf = body_shape(s).text_frame
    ps = tf.paragraphs
    six = [
        "Expand strong security foundations to the AI ecosystem",
        "Extend detection and response to bring AI into the threat universe",
        "Automate defenses to keep pace with existing and new threats",
        "Harmonize platform-level controls to ensure consistent security",
        "Adapt controls to adjust mitigations and create faster feedback loops",
        "Contextualize AI system risks in surrounding business processes",
    ]
    set_par(ps[0], "Standardized, holistic approach to integrating security and privacy into ML-powered applications")
    set_par(ps[1], "Six core elements:", level=0)
    for i, item in enumerate(six):
        set_par(ps[2 + i], item, level=1)
    for p in ps[2 + len(six):]:
        del_par(p)
    set_notes(s,
        "[flex] Corrected for rev a3: the old slide mixed sub-bullets (e.g., 'leverage secure-by-default "
        "infrastructure', 'use threat intelligence') in as if they were elements. These are Google's official "
        "six SAIF core elements; the next two SAIF slides elaborate elements 3-6 correctly.\n"
        "Source: Google — Secure AI Framework (SAIF), https://safety.google/cybersecurity-advancements/saif/ (verified 2026-08-01)")

    # --- slide 11 (idx 10): Perspective API Do Now -> rubric exercise ------ #
    s = slides[10]
    s.shapes.title.text_frame.paragraphs[0].runs[0].text = "DO NOW: Score Toxicity Like a Moderator"
    tf = body_shape(s, "Content Placeholder 2").text_frame
    ps = tf.paragraphs
    new = [
        ("Automated moderation services come and go — the judgment problem stays", 0),
        ("Score each phrase on a simple 0-3 toxicity rubric", 0),
        ("0 = civil    1 = rude    2 = abusive    3 = threatening", 1),
        ("Score these phrases, then compare with a neighbor", 0),
        ("You are stupid", 1),
        ("This relationship sucks", 1),
        ("You are acting like a jerk", 1),
        ("I know where you live", 1),
        ("Where did you disagree? Which scores depend on context?", 0),
        ("What would a false positive or a false negative cost an agency?", 0),
        ("In production, agencies pair automated filters with human review of edge cases", 0),
    ]
    for i, (text, lvl) in enumerate(new):
        set_par(ps[i], text, level=lvl)
    for p in ps[len(new):]:
        del_par(p)
    set_notes(s,
        "Replaced for rev a3: the previous exercise used the Perspective API 'TRY IT OUT' web demo. "
        "Perspective API (Google Jigsaw) is sunsetting — service officially ends December 31, 2026, usage and "
        "quota requests closed after February 2026, and no migration support is offered. This rubric-scored "
        "review needs no external service and teaches the same lesson: toxicity scoring is subjective and "
        "context-dependent, so agencies keep humans in the loop.\n"
        "Source: Google Jigsaw — Perspective API, https://www.perspectiveapi.com/ (verified 2026-08-01)")

    # --- slide 13 (idx 12): Zero Trust note correction --------------------- #
    s = slides[12]
    set_notes(s,
        "Context-Aware Access (CAA) helps manage risk by not making access a binary choice — policies can be "
        "applied flexibly to the right people, applications, and data.\n\n"
        "[flex] Corrected for rev a3: Google's BeyondCorp pioneered zero trust commercially, but the CISA Zero "
        "Trust Maturity Model v2.0 (April 2023) is not modeled on it — the ZTMM is aligned to OMB M-22-09, the "
        "federal zero-trust strategy.\n"
        "Source: CISA — Zero Trust Maturity Model, https://www.cisa.gov/zero-trust-maturity-model (verified 2026-08-01)")

    # --- final arrange: identity permutation (all 22 slides kept) ---------- #
    P.arrange(prs, list(range(len(list(prs.slides)))))
    P.save(prs, path)
    return path


def fix_appc():
    path = DECKS / "1258-AppC-Transformer-LLM-Internals.pptx"
    prs = P.open_deck(path)
    slides = list(prs.slides)

    # --- slide 14 (idx 13): encoder/decoder fork --------------------------- #
    s = slides[13]
    s.shapes.title.text_frame.paragraphs[0].runs[0].text = "The Encoder-Decoder Fork: What Actually Happened"
    tf = body_shape(s).text_frame
    ps = tf.paragraphs
    new = [
        "The 2017 Transformer has both halves: an encoder that reads, a decoder that writes",
        "BERT (2018) used only the encoder — built for understanding, not generation",
        "Decoder-only models (the GPT lineage) became the dominant paradigm for generation",
        "Encoder models persist mainly as embedding and reranker models inside search and RAG",
        "Every chat assistant your agency pilots today is a decoder-only transformer",
    ]
    for i, text in enumerate(new):
        set_par(ps[i], text)
    set_notes(s,
        "[flex] Rewritten for rev a3: the old slide framed decoder integration as a future research direction "
        "for BERT. With 2026 hindsight the fork is settled — decoder-only models won text generation while "
        "encoder models survive for understanding, embedding, and reranking tasks.\n"
        "Source: Devlin et al. (Google) — BERT: Pre-training of Deep Bidirectional Transformers, "
        "https://arxiv.org/abs/1810.04805, and Vaswani et al. — Attention Is All You Need, "
        "https://arxiv.org/abs/1706.03762 (verified 2026-08-01)")

    # --- slide 26 (idx 25): frozen weights --------------------------------- #
    s = slides[25]
    tf = body_shape(s).text_frame
    ps = tf.paragraphs
    set_par(ps[2], "Frozen weights: a deployed GPT does not learn from new data — training is finished")
    clone_par_after(ps[4], "Fresh knowledge enters through the prompt context (RAG) or retraining — not from use")
    set_notes(s,
        "[flex] Corrected for rev a3: the old 'Adaptive learning' bullet implied a deployed LLM keeps learning "
        "at inference. It does not — weights are frozen after training; new information reaches the model "
        "through the context window (RAG) or through retraining/fine-tuning.\n"
        "Source: OpenAI — Model optimization (fine-tuning) guide, https://platform.openai.com/docs/guides/fine-tuning, "
        "and Es et al. — RAGAS: Automated Evaluation of Retrieval Augmented Generation, "
        "https://arxiv.org/abs/2309.15217 (verified 2026-08-01)")

    P.arrange(prs, list(range(len(list(prs.slides)))))
    P.save(prs, path)
    return path


def fix_appd():
    path = DECKS / "1258-AppD-Classic-ML-DeepDive.pptx"
    prs = P.open_deck(path)
    slides = list(prs.slides)

    # --- slide 10 (idx 9): Word2Vec has no non-linear hidden layer --------- #
    s = slides[9]
    tf = body_shape(s).text_frame
    ps = tf.paragraphs
    set_par(ps[1], "A projection layer only — no non-linear hidden layer, which is why it trains fast")
    set_notes(s,
        "[flex] Precision fix for rev a3: Word2Vec's CBOW and Skip-gram architectures get their speed by "
        "dropping the non-linear hidden layer entirely (a projection layer only) — the paper trained "
        "high-quality vectors on 1.6 billion words in under a day.\n"
        "Source: Mikolov et al. (Google) — Efficient Estimation of Word Representations in Vector Space, "
        "https://arxiv.org/abs/1301.3781 (verified 2026-08-01)")

    # --- slide 11 (idx 10): GloVe does not use TF-IDF ---------------------- #
    s = slides[10]
    tf = body_shape(s).text_frame
    ps = tf.paragraphs
    set_par(ps[1], "Unsupervised learning — a log-bilinear model, not a neural network")
    set_par(ps[2], "Trained on large corpora: Wikipedia, Common Crawl, Twitter, Dolma (2024)")
    set_par(ps[4], "Builds a global word-word co-occurrence matrix from the whole corpus")
    set_par(ps[5], "Fits a weighted least-squares objective — no TF-IDF involved")
    set_par(ps[7], "Ratios of co-occurrence probabilities capture meaning across the corpus")
    set_notes(s,
        "[flex] Corrected for rev a3: the old slide and note claimed GloVe 'weights terms using TF-IDF' and "
        "removes stop words — neither is true. GloVe trains on aggregated global word-word co-occurrence "
        "statistics as a log-bilinear model with a weighted least-squares objective; published vectors come "
        "from Wikipedia+Gigaword, Common Crawl, Twitter, and (2024) Dolma.\n"
        "Source: Stanford NLP — GloVe: Global Vectors for Word Representation, "
        "https://nlp.stanford.edu/projects/glove/ (verified 2026-08-01)")

    P.arrange(prs, list(range(len(list(prs.slides)))))
    P.save(prs, path)
    return path


if __name__ == "__main__":
    import zipfile
    for fn, expect in ((fix_appb, 22), (fix_appc, 29), (fix_appd, 20)):
        out = fn()
        n = len(P.open_deck(out).slides._sldIdLst)
        assert n == expect, f"{out.name}: {n} slides, expected {expect}"
        assert zipfile.ZipFile(out).testzip() is None, f"{out.name}: corrupt zip"
        print(f"OK {out.name}: {n} slides, zip valid")
