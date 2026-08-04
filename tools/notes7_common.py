"""Shared helpers for the pass-7 speaker-notes rewrite (delivery-optimized notes).

Each tools/build_notes_<tag>.py script:
  1. copies the pristine deck from decks/_bak7/ back over decks/ (idempotent rebuild),
  2. rewrites every slide's speaker notes into the delivery structure
       HOOK (SAY/ASK) / KEY POINT / TIME / TRANSITION,
  3. auto-extracts preserved markers from the BACKUP notes and appends them verbatim:
       - "[flex] " prefix semantics (rewritten flex notes start with
         "[flex] (skip if behind — core message: ...)"),
       - "Source: ... (verified 2026-08-01)" citation lines (also "Sources: ..."),
       - "Prompting spine: rung Pn." lines (and spine-rung mentions),
       - author/IG flag lines (verify-before-class, re-check notes,
         "generated from the registry" lines, retitle/renumber/merge records).

Run verify() after build() to assert slide count, zip integrity, the four
required elements, device rotation, and the presence of every preserved marker.
"""
from pathlib import Path
import re
import shutil
import zipfile

from pptx import Presentation

BASE = Path(__file__).resolve().parents[1]
DECKS = BASE / "decks"
BAK = DECKS / "_bak7"

# ---------------------------------------------------------------- preserved #

LINE_START_FLAGS = (
    "[TEACH]", "AUTHOR FLAG", "Renumbered", "Retitled", "Merged in", "Repointed",
    "Corrected", "Corrects ", "Correction", "URL re-check", "Added (a4)",
    "Jogger text:", "Direction:", "Chapter starts:", "Instructor notes:",
    "Citation updated", "Two fixes", "Terminology note:", "Agenda rewritten",
    "Recap rewritten", "REFRAME", "This replaces",
)
CONTAINS_FLAGS = (
    "Generated from the activity registry",
    "prompting spine", "prompting-spine", "spine rung",
    "no source could be verified",
)
VERIFY_RE = re.compile(
    r"re-check|re-verify|re-verified|recheck|verify at teaching time|"
    r"Check the FedRAMP Marketplace", re.I)
SRC_RE = re.compile(r"Sources?: ")
SENT_SPLIT = re.compile(r"(?<=[.!?])\s+")


def extract_preserved(old_notes):
    """Return (is_flex, kept_lines) extracted verbatim from backup notes."""
    kept = []
    if not old_notes:
        return False, kept
    flex = old_notes.startswith("[flex]")
    for line in old_notes.split("\n"):
        ls = line.strip()
        if not ls:
            continue
        # 1) whole-line author/IG flags (may themselves contain a Source tail)
        low = ls.lower()
        if ls.startswith(LINE_START_FLAGS) or any(c.lower() in low for c in CONTAINS_FLAGS):
            kept.append(line.rstrip())
            continue
        # 2) verify-before-class / re-check sentences embedded in longer lines
        if VERIFY_RE.search(ls):
            for sent in SENT_SPLIT.split(ls):
                if VERIFY_RE.search(sent):
                    kept.append(sent.rstrip())
        # 3) Source citation tails (line may start with other prose)
        m = SRC_RE.search(ls)
        if m and "(verified 2026-08-01)" in ls[m.start():]:
            seg = ls[m.start():].rstrip()
            if not any(seg in k for k in kept):
                kept.append(seg)
    return flex, kept


# ------------------------------------------------------------------- build #

def build(deck_name, spec):
    """Rebuild decks/<deck_name> from _bak7 and rewrite all notes.

    spec: dict 1-based slide index -> dict with keys
      hook (starts 'ASK:' / 'SAY (scenario):' / 'SAY (fact):'),
      key, time, transition, and optional flex_msg (required on [flex] slides).
      Set force_flex=True on activity slides whose activity is tagged [flex]
      on the chapter's activities slide even though the old notes lacked the
      prefix (e.g., Lab 4.2, Lab 5.1, DO NOW 5.C).
    """
    src = BAK / deck_name
    dst = DECKS / deck_name
    shutil.copy(src, dst)                      # idempotent: pristine base
    prs = Presentation(dst)
    slides = list(prs.slides)
    assert len(slides) == len(spec), f"{deck_name}: {len(slides)} slides vs {len(spec)} specs"
    for i, slide in enumerate(slides, 1):
        old = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        flex, kept = extract_preserved(old)
        flex = flex or bool(spec[i].get("force_flex"))
        s = spec[i]
        lines = []
        if flex:
            assert s.get("flex_msg"), f"{deck_name} slide {i}: [flex] slide needs flex_msg"
            lines.append(f"[flex] (skip if behind — core message: {s['flex_msg']})")
        lines.append(s["hook"])
        lines.append(f"KEY POINT: {s['key']}")
        lines.append(f"TIME: {s['time']}")
        lines.append(f"TRANSITION: {s['transition']}")
        if kept:
            lines.append("")
            lines.extend(kept)
        slide.notes_slide.notes_text_frame.text = "\n".join(lines)
    prs.save(dst)
    return dst


# ------------------------------------------------------------------ verify #

DEVICES = ("ASK", "SAY (scenario)", "SAY (fact)")


def _device_of(notes_text):
    for line in notes_text.split("\n"):
        ls = line.strip()
        if not ls or ls.startswith("[flex]"):
            continue
        for d in DEVICES:
            if ls.startswith(d + ":"):
                return d
        return None
    return None


def verify(deck_name, expected_count):
    """Assert slide count, zip integrity, 4 elements, device rotation, markers."""
    dst = DECKS / deck_name
    problems = []

    prs = Presentation(dst)
    slides = list(prs.slides)
    if len(slides) != expected_count:
        problems.append(f"slide count {len(slides)} != {expected_count}")

    with zipfile.ZipFile(dst) as z:
        bad = z.testzip()
        if bad is not None:
            problems.append(f"zip corrupt at {bad}")

    bak_slides = list(Presentation(BAK / deck_name).slides)
    prev_dev, prev_idx = None, None
    for i, slide in enumerate(slides, 1):
        if not slide.has_notes_slide:
            problems.append(f"s{i}: no notes slide")
            continue
        t = slide.notes_slide.notes_text_frame.text
        for elem in ("KEY POINT:", "TIME:", "TRANSITION:"):
            if elem not in t:
                problems.append(f"s{i}: missing {elem}")
        dev = _device_of(t)
        if dev is None:
            problems.append(f"s{i}: missing SAY/ASK hook")
        elif dev == prev_dev:
            problems.append(f"s{i}: device {dev!r} same as s{prev_idx}")
        prev_dev, prev_idx = dev, i

        old = bak_slides[i - 1].notes_slide.notes_text_frame.text \
            if bak_slides[i - 1].has_notes_slide else ""
        flex, kept = extract_preserved(old)
        if flex and not t.startswith("[flex] (skip if behind"):
            problems.append(f"s{i}: lost [flex] prefix")
        for k in kept:
            if k not in t:
                problems.append(f"s{i}: lost preserved marker: {k[:70]}...")
    return problems
