"""Bring the decks into line with the a4 activity architecture.

Three operations per chapter deck:

  1. RETITLE the existing dedicated launcher slides whose id or title drifted
     (several still carry a1/a2 wording, e.g. "Lab 5.2: Data Minimization
     Checklist" for what is now the PII lab).
  2. INSERT one "Activities in This Chapter" slide near the front, generated
     from content_activities.py, so the room can see the demo, the four Do Nows
     and the labs with their timings.
  3. APPEND a launcher slide for each lab that does not already have one.

Slide numbers are then written back into the registry and the lab manual, so the
manual's "Where this fits" line points at real slides.

Run:  python tools/build_a4_slides.py
"""
from pathlib import Path

import yaml

import a4_common as A
import content_activities as C
import pptx_tools as P

ROOT = Path(__file__).resolve().parent.parent

# existing dedicated launcher slides that need their wording brought up to date:
# (chapter, 1-based slide, new title, new bullets)
RETITLE = {
    "Ch01": [(23, "Lab 1.1: Exploring the Federal AI Use Case Inventory",
              ["Load the real OMB inventory — 1,500+ federal AI use cases",
               "Group by agency, development stage and high-impact flag",
               "Write three findings, each with the code that produced it",
               "No AI needed — this is the data literacy underneath it. See Lab Manual, Lab 1.1"])],
    "Ch05": [(21, "Lab 5.1: Adversarial Examples and Model Robustness",
              ["Train a small classifier, then break it on purpose",
               "Find the smallest perturbation that flips a prediction",
               "Apply one defence and re-measure",
               "See Lab Manual, Lab 5.1"]),
             (32, "Lab 5.2: Detect and Mask PII in Citizen Records",
              ["Run a detector over synthetic records and real 311 free text",
               "Find three items it misses by reading rows yourself",
               "Mask without destroying analytic value; report residual risk",
               "See Lab Manual, Lab 5.2"])],
    "Ch06": [(14, "DO NOW 6.B: Propose the Schema",
              ["Describe a new collection need to an assistant",
               "Ask for field, type, required, allowed values, example",
               "Ask which three defects will appear anyway",
               "Prompting spine rung P6. See Lab Manual, DO NOW 6.B"]),
             (15, "Lab 6.1: Data Quality Assessment of Chicago 311",
              ["Assess real operational data against all six dimensions",
               "Quantify each defect — null rates, duplicates, invalid dates",
               "Build a scorecard and a go / no-go recommendation",
               "See Lab Manual, Lab 6.1"])],
    "Ch08": [(39, "Lab 8.1: Visualization and Reporting from EPA Data",
              ["Rank counties by unhealthy days; chart the top 15",
               "Apply the guidelines: sorted, labelled, honest scale",
               "Assemble a one-page briefing with a recommendation",
               "See Lab Manual, Lab 8.1"])],
}

# labs that have no launcher slide yet
NEEDS_LAUNCHER = {
    "Ch00": ["Lab 0.1"], "Ch01": ["Lab 1.2"], "Ch02": ["Lab 2.1"],
    "Ch04": ["Lab 4.2"], "Ch06": ["Lab 6.2"], "Ch08": ["Lab 8.2"], "Ch09": ["Lab 9.1"],
}


def by_id(aid):
    return next(a for a in C.ALL if a["id"] == aid)


def activities_slide(ch, items):
    """The one overview slide: demo, four Do Nows, labs, with timings."""
    demo = [a for a in items if a["kind"] == "demo"]
    dns = [a for a in items if a["kind"] == "donow"]
    labs = [a for a in items if a["kind"] == "lab"]
    bullets = []
    if demo:
        d = demo[0]
        bullets.append(f"Demo (instructor): {d['title']} — {d['minutes']} min")
    for a in dns:
        tag = " [flex]" if a["flex"] else ""
        bullets.append(f"{a['id']}: {a['title']} — {a['minutes']} min{tag}")
    for a in labs:
        tag = " [flex]" if a["flex"] else ""
        bullets.append(f"{a['id']}: {a['title']} — {a['minutes']} min{tag}")
    bullets.append("Full instructions for all of these are in your Lab Manual")
    notes = (f"[TEACH] Activities overview for {ch}. Generated from the activity registry "
             f"(registry/activities.yaml) — do not hand-edit; edit tools/content_activities.py "
             f"and re-run tools/build_a4_slides.py. Chapter activity budget: "
             f"{sum(a['minutes'] for a in items)} min "
             f"({sum(a['minutes'] for a in items if not a['flex'])} core).")
    return ("Activities in This Chapter", bullets, notes)


def launcher_slide(a):
    steps = a["steps"][:3]
    bullets = [s.replace("**", "").split(". ", 1)[-1] if s[:2].isdigit() else s.replace("**", "")
               for s in steps]
    bullets = [b if len(b) < 105 else b[:102] + "..." for b in bullets]
    bullets.append(f"{a['minutes']} minutes. See Lab Manual, {a['id']}")
    notes = f"[TEACH] Launcher for {a['id']}. {a['scenario']}"
    if a["data"]:
        notes += " Data: " + ", ".join(a["data"]) + "."
    if a["needs_key"]:
        notes += " Uses the OpenAI key from .env."
    return (f"{a['id']}: {a['title']}", bullets, notes)


def process(ch, items):
    deck = A.deck(C.DECK_OF[ch])
    prs = P.open_deck(deck)
    A.clean_orphan_slides(prs)
    P.save(prs, deck)
    prs = P.open_deck(deck)
    slides = list(prs.slides)
    n0 = len(slides)
    changed = []

    # 1. retitle drifted launchers
    for idx, new_title, new_bullets in RETITLE.get(ch, []):
        s = slides[idx - 1]
        body = A._body_shape(s)
        title_sh = None
        for sh in s.shapes:
            if sh.has_text_frame and sh is not body:
                title_sh = sh
        if title_sh is not None:
            tf = title_sh.text_frame
            if tf.paragraphs and tf.paragraphs[0].runs:
                tf.paragraphs[0].runs[0].text = new_title
                for r in tf.paragraphs[0].runs[1:]:
                    r.text = ""
        if body is not None:
            tf = body.text_frame
            donor = next((p for p in tf.paragraphs if p.runs), None)
            for p in list(tf.paragraphs):
                for r in p.runs:
                    r.text = ""
            if donor and donor.runs:
                donor.runs[0].text = new_bullets[0]
            for b in new_bullets[1:]:
                A.add_bullet(s, b)
        A.append_notes(s, f"Retitled in a4 to match the activity registry ({new_title}).")
        changed.append(f"retitled s{idx}")

    # 2 + 3. append the overview and any missing launchers, then place them
    new_items = [activities_slide(ch, items)]
    for aid in NEEDS_LAUNCHER.get(ch, []):
        new_items.append(launcher_slide(by_id(aid)))
    for t, b, n in new_items:
        P.append_content(prs, t, b, n)

    appended = list(range(n0, n0 + len(new_items)))
    overview, launchers = appended[0], appended[1:]
    # overview goes third (after title + objectives); launchers just before the last slide
    order = [0, 1, overview] + [i for i in range(2, n0 - 1)] + launchers + [n0 - 1]
    assert sorted(order) == list(range(n0 + len(new_items))), f"{ch}: order not a permutation"
    P.arrange(prs, order)
    P.save(prs, deck)

    n = A.verify(deck)
    changed.append(f"+{len(new_items)} slides")
    print(f"  {ch}: {', '.join(changed)} -> {n} slides")
    return n


def write_slide_refs():
    """Where each activity id appears, across EVERY deck including appendices.

    Appendices matter: AppD points back at Ch01's labs, so a chapter-only scan
    reports drift that is not really drift.
    """
    from pptx import Presentation
    texts = {}
    for p in sorted(A.DECKS.glob("*.pptx")):
        stem = p.stem.replace("1258-", "")
        texts[stem] = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
                       for s in Presentation(str(p)).slides]
    refs = {}
    for a in C.ALL:
        v = a["id"].lower()
        refs[a["id"]] = [f"{stem}:{i}" for stem, slides in texts.items()
                         for i, t in enumerate(slides, 1) if v in t.lower()]
    return refs


if __name__ == "__main__":
    print("a4 pass 4 — align decks with the activity architecture")
    total = 0
    for ch, items in C.CHAPTERS:
        total += process(ch, items)
    print(f"\n  core total: {total} slides")

    refs = write_slide_refs()
    # push slide refs into the generated registry
    reg_path = ROOT / "registry" / "activities.yaml"
    doc = yaml.safe_load(reg_path.read_text())
    for row in doc["activities"]:
        row["slide_ref"] = [f"{row['deck']}:{n}" for n in refs.get(row["id"], [])]
    reg_path.write_text(yaml.safe_dump(doc, sort_keys=False, allow_unicode=True, width=100))
    missing = [i for i, v in refs.items() if not v]
    print(f"  slide_ref written for {len(refs) - len(missing)}/{len(refs)} activities")
    if missing:
        print(f"  no slide reference yet: {', '.join(missing)}")

    # regenerate the manual with real slide numbers
    import build_a4_activities as B
    hint = {k: (", ".join(str(n) for n in v) if v else "see chapter deck")
            for k, v in refs.items()}
    B.build_manual(hint)
    print("  lab manual regenerated with slide references")
