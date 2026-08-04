"""Final consistency pass on the enriched a3 decks.

(a) Renumber stale chapter/lab labels the enrichment passes left in place
    (they operated under a no-renumber rule; MOVES.md + registry/labs.yaml
    are now the authority):
    - Ch05 s1 banner/notes "Chapter 3 / Day 1 2:15pm / jogger EDA" -> Ch 5,
      Day 2 PM, jogger "AI Security"; lab launchers Lab 3.1/3.2/3.4 -> 5.1/5.2/5.4
      (no Lab 3.3 launcher exists in the deck — see report)
    - Ch06 s1 "Chapter 5" -> "Chapter 6", Day 2 PM
    - Ch07 s1 "Chapter 6" -> "Chapter 7", Day 3 AM
    - Ch08 s1 "Chapter 4" -> "Chapter 8", Day 3 PM; s46 "Lab 7.1" -> "Lab 8.1"
(b) Ch02 residual fabricated/unverified stats, per
    1258/research/refs-ch02-gov-apps-platforms.md and
    4703/research/refs-appc-case-studies.md (slides 7, 16, 17, 19, 24, 25).
(c) Ch05 s41 Do Now URL: tensorflow.org/responsible_ai re-verified live
    2026-08-01 (HTTP 200, "Responsible AI Toolkit") — kept, noted.

Every run-level edit preserves formatting; every touched slide gets a speaker
note documenting the change and its source. Idempotent: a match that is
already fixed is skipped (asserts the new text is present instead).
"""
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pptx_tools as P

DECKS = Path(__file__).resolve().parents[1] / "decks"
V = "(verified 2026-08-01)"


# ---------------------------------------------------------------- helpers
def move_orphan_slide_parts(prs):
    """Rename slide parts that are reachable but NOT in sldIdLst (dead orphans).
    Copied from build_enrich_ch01.py — defensive: none of the decks touched
    here currently have orphans, but saving with one would duplicate zip
    entries. Must run before ANY access of prs.slides."""
    import re
    from pptx.oxml.ns import qn
    from pptx.opc.packuri import PackURI

    sldIdLst = prs.part._element.get_or_add_sldIdLst()
    keep = {id(prs.part.rels[e.get(qn("r:id"))].target_part) for e in sldIdLst}
    pat = re.compile(r"/ppt/slides/slide(\d+)\.xml$")
    parts = [p for p in prs.part.package.iter_parts() if pat.match(str(p.partname))]
    used = {int(pat.match(str(p.partname)).group(1)) for p in parts}
    nxt = max(used) + 1
    moved = 0
    for p in parts:
        if id(p) not in keep:
            p.partname = PackURI(f"/ppt/slides/slide{nxt}.xml")
            nxt += 1
            moved += 1
    return moved


def _paras(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            yield para


def edit_para(slide, match, new_text):
    """Replace a whole paragraph's text, preserving the first run's formatting.
    Idempotent: if `match` is gone, asserts `new_text` (or its head) is there."""
    hits = 0
    for para in _paras(slide):
        if match in para.text:
            if para.runs:
                para.runs[0].text = new_text
                for run in para.runs[1:]:
                    run.text = ""
            hits += 1
    if hits == 0:
        assert any(new_text[:30] in para.text for para in _paras(slide)), \
            f"edit_para {match!r}: 0 hits and replacement not present"
        return False
    assert hits == 1, f"edit_para {match!r}: {hits} hits"
    return True


def del_para(slide, match):
    """Delete the single paragraph containing `match` (idempotent)."""
    hits = 0
    for para in _paras(slide):
        if match in para.text:
            para._element.getparent().remove(para._element)
            hits += 1
    assert hits <= 1, f"del_para {match!r}: {hits} hits"
    return hits == 1


def fix_notes(slide, old, new):
    """String-replace inside the notes text frame (idempotent)."""
    tf = slide.notes_slide.notes_text_frame
    if old in tf.text:
        tf.text = tf.text.replace(old, new)
        return True
    assert new in tf.text, f"fix_notes {old!r}: neither old nor new present"
    return False


def append_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    if text in tf.text:
        return False
    tf.text = (tf.text.rstrip() + "\n\n" + text) if tf.text.strip() else text
    return True


def get_deck(name, expect):
    prs = Presentation(DECKS / name)
    n = move_orphan_slide_parts(prs)
    if n:
        print(f"  {name}: moved {n} orphan slide parts")
    assert len(list(prs.slides)) == expect, f"{name}: expected {expect} slides"
    return prs, list(prs.slides)


# ---------------------------------------------------------------- Ch05
def fix_ch05():
    prs, s = get_deck("1258-Ch05-Security-Risks-Responsible-AI.pptx", 58)

    # s1: chapter banner + jogger/day notes (a3 renumber: old Ch03 -> new Ch05)
    edit_para(s[0], "Chapter 3", "Chapter 5")
    fix_notes(s[0], "Jogger text: Exploratory Data Analysis", "Jogger text: AI Security")
    fix_notes(s[0], "Chapter starts: Day 1 at 2:15pm", "Chapter starts: Day 2 (PM session)")
    append_notes(s[0],
        "Renumbered for rev a3: this deck is the old Chapter 3 (security) and now opens "
        "Chapter 5 on Day 2 PM per MOVES.md; the jogger previously read 'Exploratory Data "
        "Analysis' (inherited from the old data chapter).")

    # lab launchers: old Lab 3.x -> new Lab 5.x (registry/labs.yaml)
    for idx, old, new in ((20, "Lab 3.1", "Lab 5.1"),
                          (31, "Lab 3.2", "Lab 5.2"),
                          (46, "Lab 3.4", "Lab 5.4")):
        hits = 0
        for para in list(_paras(s[idx])):
            if old in para.text:
                for run in para.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        hits += 1
        assert hits == 2 or (hits == 0 and new in "\n".join(
            p.text for p in _paras(s[idx]))), f"s{idx+1}: {old} hits={hits}"
        append_notes(s[idx],
            f"Renumbered for rev a3: {old} is now {new} per MOVES.md and "
            "registry/labs.yaml (title unchanged).")

    # s41 Do Now: tensorflow.org/responsible_ai re-verified live (was redirect-looping)
    append_notes(s[40],
        "URL re-check (a3): https://www.tensorflow.org/responsible_ai previously "
        "redirect-looped; as of 2026-08-01 it serves HTTP 200 as the 'Responsible AI "
        "Toolkit' page (TensorFlow's ecosystem of responsible-AI tools), which matches "
        "this Do Now's mapping exercise — URL kept. Fallback if it breaks again: the "
        "NIST AI RMF hub, https://www.nist.gov/itl/ai-risk-management-framework.")

    P.save(prs, DECKS / "1258-Ch05-Security-Risks-Responsible-AI.pptx")
    print("  saved Ch05")


# ---------------------------------------------------------------- Ch06
def fix_ch06():
    prs, s = get_deck("1258-Ch06-Data-for-AI.pptx", 52)
    edit_para(s[0], "Chapter 5", "Chapter 6")
    fix_notes(s[0], "Chapter starts: Day 1 at 2:15pm", "Chapter starts: Day 2 (PM session)")
    append_notes(s[0],
        "Renumbered for rev a3: this deck is the old Chapter 5 (data) and now opens "
        "Chapter 6 on Day 2 PM per MOVES.md. Jogger 'Exploratory Data Analysis' is "
        "correct for this chapter and was kept.")
    P.save(prs, DECKS / "1258-Ch06-Data-for-AI.pptx")
    print("  saved Ch06")


# ---------------------------------------------------------------- Ch07
def fix_ch07():
    prs, s = get_deck("1258-Ch07-Building-with-LLMs.pptx", 56)
    edit_para(s[0], "Chapter 6", "Chapter 7")
    append_notes(s[0],
        "Chapter starts: Day 3 (AM session).\n"
        "Renumbered for rev a3: the old Chapter 6 was split into Ch04 (prompt "
        "engineering) and this deck, now Chapter 7, per MOVES.md.")
    P.save(prs, DECKS / "1258-Ch07-Building-with-LLMs.pptx")
    print("  saved Ch07")


# ---------------------------------------------------------------- Ch08
def fix_ch08():
    prs, s = get_deck("1258-Ch08-Operations-Reporting.pptx", 52)
    edit_para(s[0], "Chapter 4", "Chapter 8")
    fix_notes(s[0], "Jogger text: Exploratory Data Analysis", "Jogger text: AI Operations")
    fix_notes(s[0], "Chapter starts: Day 1 at 2:15pm", "Chapter starts: Day 3 (PM session)")
    append_notes(s[0],
        "Renumbered for rev a3: this deck merges the old Ch04 (operations) and old Ch07 "
        "(reporting/visualization) into Chapter 8 on Day 3 PM per MOVES.md; the jogger "
        "previously read 'Exploratory Data Analysis' (inherited from the old data chapter).")

    # s46: old Lab 7.1 (visualization) -> new Lab 8.1
    hits = 0
    for para in list(_paras(s[45])):
        if "Lab 7.1" in para.text:
            for run in para.runs:
                if "Lab 7.1" in run.text:
                    run.text = run.text.replace("Lab 7.1", "Lab 8.1")
                    hits += 1
    assert hits == 2 or (hits == 0 and "Lab 8.1" in "\n".join(
        p.text for p in _paras(s[45]))), f"s46: Lab 7.1 hits={hits}"
    append_notes(s[45],
        "Renumbered for rev a3: old Lab 7.1 (Visualization Using Python) is now Lab 8.1 "
        "per MOVES.md and registry/labs.yaml (title unchanged; it is a flex lab).")
    P.save(prs, DECKS / "1258-Ch08-Operations-Reporting.pptx")
    print("  saved Ch08")


# ---------------------------------------------------------------- Ch02
def fix_ch02():
    prs, s = get_deck("1258-Ch02-Applications-Public-Data.pptx", 52)

    # --- s7: Google retinopathy "90% accuracy" + radiology "-30%" ---
    edit_para(s[6], "90% accuracy",
              "Example: Google's deep-learning model detected diabetic retinopathy with "
              "~90% sensitivity and ~98% specificity in validation (Gulshan et al., JAMA 2016)*")
    edit_para(s[6], "reduced radiology reporting time by 30%",
              "Example: AI-assisted reporting shortens radiology turnaround time")
    edit_para(s[6], "health.google/ai-retinopathy",
              "*Gulshan et al., JAMA 2016 — "
              "https://jamanetwork.com/journals/jama/fullarticle/2588763")
    append_notes(s[6],
        "Corrections (a3): '90% accuracy' was a loose paraphrase — Gulshan et al. (JAMA "
        "2016) report sensitivity/specificity pairs, not a single accuracy figure "
        "(~97.5%/93.4% high-sensitivity point; ~90.3%/98.1% high-specificity point on "
        "EyePACS-1); the slide now cites the paper instead of the health.google marketing "
        "link. 'AI reduced radiology reporting time by 30%' could not be verified anywhere "
        "and was softened to a qualitative claim.\n"
        "Source: Gulshan, Peng, et al. — 'Development and Validation of a Deep Learning "
        "Algorithm for Detection of Diabetic Retinopathy', JAMA 316(22), 2016, "
        "https://jamanetwork.com/journals/jama/fullarticle/2588763 " + V)

    # --- s16: Siemens "20%" + Pittsburgh surtrac.net citation ---
    edit_para(s[15], "Siemens AI predicted train track failures",
              "Example: rail operators use AI to predict track and equipment failures "
              "(vendor-reported)")
    edit_para(s[15], "AI-driven traffic lights",
              "Example: Pittsburgh's SURTRAC adaptive signals cut wait times ~40% and "
              "travel times ~25% in field tests (CMU, ICAPS 2013)")
    edit_para(s[15], "siemens.com/smart-maintenance",
              "*Vendor-reported; no independent source located")
    edit_para(s[15], "surtrac.net",
              "**Smith et al. (CMU), ICAPS 2013 — https://aaai.org/papers/"
              "00434-13594-smart-urban-signal-networks-initial-application-of-the-"
              "surtrac-adaptive-traffic-signal-control-system/")
    append_notes(s[15],
        "Corrections (a3): the Siemens '20% maintenance cost reduction' is an unverified "
        "vendor claim (siemens.com/smart-maintenance is not a real citation URL) — figure "
        "dropped and labeled vendor-reported. The Pittsburgh '40% wait time' figure IS "
        "verified: CMU's SURTRAC pilot (nine East Liberty intersections) reported ~40% "
        "wait-time, ~25% travel-time, and ~20% emissions reductions in the peer-reviewed "
        "ICAPS 2013 paper, which now replaces the surtrac.net marketing link.\n"
        "Source: Smith, Barlow, Xie & Rubinstein (Carnegie Mellon) — 'Smart Urban Signal "
        "Networks: Initial Application of the SURTRAC Adaptive Traffic Signal Control "
        "System', ICAPS 2013, https://aaai.org/papers/00434-13594-smart-urban-signal-"
        "networks-initial-application-of-the-surtrac-adaptive-traffic-signal-control-system/ " + V)

    # --- s17: Barcelona smart bins (unverifiable) ---
    edit_para(s[16], "smart bins",
              "Example: cities apply AI to route waste collection and manage transit demand")
    del_para(s[16], "barcelona.cat/smartcity")
    append_notes(s[16],
        "Correction (a3): 'Barcelona's smart bins reduced waste collection costs' could "
        "not be verified in any official or independent source (barcelona.cat/smartcity "
        "is not a real citation URL) and was replaced with a qualitative smart-city claim. "
        "Do not quote a figure here unless you can cite it. (4703 research pass, gaps list.)")

    # --- s19: FEMA NLP claim (no FEMA-primary source) ---
    edit_para(s[18], "FEMA uses NLP-based systems",
              "DHS components field AI for disaster response — the DHS AI Use Case "
              "Inventory lists current, named use cases (FEMA, TSA, USCIS, CISA)")
    edit_para(s[18], "Identifies critical areas",
              "Check the inventory for what is deployed today: "
              "https://www.dhs.gov/ai/use-case-inventory")
    append_notes(s[18],
        "Correction (a3): 'FEMA uses NLP-based systems to analyze social media posts and "
        "emergency call data' has no FEMA-primary source (two research passes found none) "
        "and was replaced with a pointer to the DHS AI Use Case Inventory — the same "
        "authoritative source the 'What Is Your Agency Doing?' slide uses later in this "
        "deck. Verified alternative disaster-AI example if you want a named case: Google's "
        "AI flood forecasting (Nearing et al., Nature 2024).\n"
        "Source: U.S. DHS — AI Use Case Inventory, https://www.dhs.gov/ai/use-case-inventory " + V)

    # --- s24: UK Home Office "60%" ---
    edit_para(s[23], "UK Home Office cut visa processing times by 60%",
              "Example: UK Home Office halted its visa-streaming algorithm in 2020 after "
              "a legal challenge over nationality bias (Foxglove/JCWI)")
    append_notes(s[23],
        "Correction (a3): 'cut visa processing times by 60%' is unverified. The documented "
        "Home Office algorithm story is the opposite lesson: a 2015–2020 visa-streaming "
        "algorithm weighted applications by nationality into green/amber/red streams and "
        "was halted in August 2020 after the Foxglove/JCWI judicial-review challenge — a "
        "governance cautionary tale (pairs with the PredPol slide later in this deck).\n"
        "Source: Digital Freedom Fund — 'UK Home Office Visa Application Streaming "
        "Algorithm', https://digitalfreedomfund.org/case-studies/uk-home-office-visa-"
        "application-streaming-algorithm/ " + V)

    # --- s25: Ask Jamie "90%" + Canada welfare AI ---
    edit_para(s[24], "Ask Jamie",
              "Example: Singapore's 'Ask Jamie' agency chatbots are being retired in "
              "favor of GovTech's LLM-based VICA platform*")
    del_para(s[24], "Canada used AI")
    edit_para(s[24], "govtech.gov.sg/askjamie",
              "*GovTech VICA — https://www.tech.gov.sg/products-and-services/vica/")
    del_para(s[24], "canada.ca/ai-policy")
    append_notes(s[24],
        "Corrections (a3): Ask Jamie's '90% success' figure is unverified and the example "
        "was dated — Singapore began retiring Ask Jamie in 2023 in favor of VICA, "
        "GovTech's next-generation LLM-based chatbot platform; the slide now says so and "
        "cites the official VICA page (govtech.gov.sg/askjamie was the wrong domain). "
        "'Canada used AI to adjust welfare policies' had no source anywhere (recommended "
        "for deletion in both research passes) and was deleted, along with the fake "
        "canada.ca/ai-policy footnote.\n"
        "Source: GovTech Singapore — 'VICA (Virtual Intelligent Chat Assistant)', "
        "https://www.tech.gov.sg/products-and-services/vica/ " + V)

    P.save(prs, DECKS / "1258-Ch02-Applications-Public-Data.pptx")
    print("  saved Ch02")


def main():
    fix_ch05()
    fix_ch06()
    fix_ch07()
    fix_ch08()
    fix_ch02()
    print("done")


if __name__ == "__main__":
    main()
