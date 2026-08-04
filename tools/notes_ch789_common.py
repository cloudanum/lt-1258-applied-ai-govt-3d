"""
notes_ch789_common.py — shared helpers for the Ch07/Ch08/Ch09 delivery-notes
rewrite (build_notes_ch07/08/09.py). Uniquely named so parallel notes passes
on other chapters do not collide with it.

Each build_notes_<tag>.py:
  1. restores its deck from decks/_bak7/ (so every run is idempotent),
  2. rewrites EVERY slide's speaker notes into the 4-element delivery
     structure (SAY/ASK hook, KEY POINT, TIME, TRANSITION),
  3. appends preserved markers auto-extracted verbatim from the backup notes,
  4. verifies structure, markers, device rotation, count, and zip integrity.

Preserved-marker policy (extracted line-by-line from the backup notes):
  - "Source: / Sources: ... (verified 2026-08-01)" citations — kept exactly;
    when embedded mid-paragraph, the citation segment from "Source(s):" to
    end of line is kept exactly.
  - "Prompting spine: rung Pn." lines.
  - "[TEACH]" registry-generated / launcher flag lines.
  - "Rev a3:" / "Renumbered for rev a3:" / "Merged in a4:" / "Retitled in a4"
    author/IG flag lines.
  - production metadata on title slides (Jogger text / Direction /
    Chapter starts / Instructor notes: <content>).
  - verify-before-class / recheck flags.
  - fabrication-guard cautions ("could not be verified", "illustrative
    scenario, not a measured case study").
  - [flex] semantics: slides whose backup notes open with "[flex]" get new
    notes opening "[flex] (skip if behind — core message: ...)".
"""
from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation

# Whole-line preserved markers (line must match to be kept verbatim).
PRESERVE_PATTERNS = [
    r"^Sources?:",
    r"^Prompting spine:",
    r"^\[TEACH\]",
    r"^Rev a3:",
    r"^Renumbered for rev a3:",
    r"^Merged in a4:",
    r"^Retitled in a4",
    r"^Jogger text:",
    r"^Direction:",
    r"^Chapter starts:",
    r"^Instructor notes:\s*\S",
    r"(?i)verify[- ]before[- ]class",
    r"(?i)\brecheck\b",
    r"could not be verified",
    r"illustrative scenario, not a measured case study",
]
_RX = [re.compile(p) for p in PRESERVE_PATTERNS]

# Citation segments embedded mid-paragraph in older notes: keep the citation
# itself ("Source: ... (verified 2026-08-01)") exactly, from token to EOL.
_SOURCE_RX = re.compile(r"\bSources?:\s")

HOOK_LABEL = {
    "ask": "ASK",        # question to the room
    "hands": "HANDS",    # show-of-hands poll
    "scenario": "SAY",   # relatable gov-work scenario
    "fact": "SAY",       # surprising fact
    "do": "SAY",         # action / demo cue
}


def extract_preserved(notes: str) -> list[str]:
    """Return the preserved-marker segments of a backup note, verbatim."""
    out = []
    for line in (notes or "").split("\n"):
        line = line.rstrip()
        if not line.strip():
            continue
        if any(rx.search(line) for rx in _RX):
            out.append(line)
            continue
        m = _SOURCE_RX.search(line)
        if m:
            out.append(line[m.start():])
    # de-duplicate, keep order
    seen, dedup = set(), []
    for k in out:
        if k not in seen:
            seen.add(k)
            dedup.append(k)
    return dedup


def compose(entry: tuple, preserved: list[str]) -> str:
    """Build the final note text for one slide.

    entry = (device, hook, key_point, time_cue, transition, flex_core|None)
    """
    device, hook, key, time_cue, transition, flex_core = entry
    lines = []
    if flex_core:
        lines.append(f"[flex] (skip if behind — core message: {flex_core})")
    lines.append(f"{HOOK_LABEL[device]}: {hook}")
    lines.append(f"KEY POINT: {key}")
    lines.append(f"TIME: {time_cue}")
    lines.append(f"TRANSITION: {transition}")
    if preserved:
        lines.append("")
        lines.extend(preserved)
    return "\n".join(lines)


def rebuild(deck_path: Path, notes: dict[int, tuple]) -> Path:
    """Restore deck from _bak7 and rewrite every slide's notes. Idempotent."""
    bak = deck_path.parent / "_bak7" / deck_path.name
    assert bak.exists(), f"missing backup {bak}"
    shutil.copy(bak, deck_path)
    bak_prs = Presentation(bak)
    prs = Presentation(deck_path)
    slides = list(prs.slides)
    bak_slides = list(bak_prs.slides)
    assert len(slides) == len(bak_slides) == len(notes), (
        f"slide/notes count mismatch: deck={len(slides)} bak={len(bak_slides)} notes={len(notes)}"
    )
    for i, slide in enumerate(slides):
        bak_notes = (
            bak_slides[i].notes_slide.notes_text_frame.text
            if bak_slides[i].has_notes_slide
            else ""
        )
        preserved = extract_preserved(bak_notes)
        slide.notes_slide.notes_text_frame.text = compose(notes[i + 1], preserved)
    prs.save(deck_path)
    return deck_path


def verify(deck_path: Path, notes: dict[int, tuple], tag: str) -> dict:
    """Scripted assertions on the rebuilt deck. Returns stats for the report."""
    bak = deck_path.parent / "_bak7" / deck_path.name
    prs = Presentation(deck_path)
    bak_prs = Presentation(bak)
    slides = list(prs.slides)
    bak_slides = list(bak_prs.slides)

    # 1. slide count unchanged
    assert len(slides) == len(bak_slides), "slide count changed"

    # 2. zip integrity
    assert zipfile.ZipFile(deck_path).testzip() is None, "zipfile.testzip failed"

    stats = {"slides": len(slides), "preserved_total": 0, "by_kind": {}}
    prev_device = None
    for i, slide in enumerate(slides):
        n = slide.notes_slide.notes_text_frame.text
        entry = notes[i + 1]
        device, _, _, _, _, flex_core = entry

        # 3. the 4 delivery elements are present
        body_lines = n.split("\n")
        first = body_lines[1] if flex_core else body_lines[0]
        assert re.match(r"^(SAY|ASK|HANDS): ", first), f"slide {i+1}: bad hook line"
        assert "KEY POINT: " in n, f"slide {i+1}: missing KEY POINT"
        assert re.search(r"^TIME: .*(min|~)", n, re.M), f"slide {i+1}: missing TIME cue"
        assert "TRANSITION: " in n, f"slide {i+1}: missing TRANSITION"

        # 4. flex semantics (both directions)
        bak_text = (
            bak_slides[i].notes_slide.notes_text_frame.text
            if bak_slides[i].has_notes_slide
            else ""
        )
        if bak_text.lstrip().startswith("[flex]"):
            assert flex_core, f"slide {i+1}: backup was [flex] but entry lacks flex_core"
        if flex_core:
            assert n.startswith("[flex] (skip if behind — core message: "), (
                f"slide {i+1}: flex note must open with [flex] marker"
            )

        # 5. every preserved marker from the backup is present verbatim
        preserved = extract_preserved(bak_text)
        for line in preserved:
            assert line in n, f"slide {i+1}: lost preserved line: {line[:60]}..."
        stats["preserved_total"] += len(preserved)
        for line in preserved:
            kind = next((p for p, rx in zip(PRESERVE_PATTERNS, _RX) if rx.search(line)),
                        "mid-line Source(s)")
            stats["by_kind"][kind] = stats["by_kind"].get(kind, 0) + 1

        # 6. delivery block shape (4 lines, or 5 with the flex opener)
        delivery_len = 5 if flex_core else 4
        assert body_lines[delivery_len - 1].startswith("TRANSITION: "), (
            f"slide {i+1}: unexpected delivery-block length"
        )

        # 7. no two adjacent slides with the same hook device
        assert device != prev_device, f"slide {i+1}: same device as previous slide"
        prev_device = device

    print(f"[{tag}] OK: {stats['slides']} slides rewritten, "
          f"{stats['preserved_total']} preserved markers, zip OK")
    return stats


def show_samples(deck_path: Path, slide_nos: list[int], tag: str) -> None:
    prs = Presentation(deck_path)
    slides = list(prs.slides)
    for no in slide_nos:
        print(f"\n----- {tag} sample: slide {no} -----")
        print(slides[no - 1].notes_slide.notes_text_frame.text)
