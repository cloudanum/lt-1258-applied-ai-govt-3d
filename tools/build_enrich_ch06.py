"""Enrich 1258-Ch06-Data-for-AI.pptx in place (a3 revision).

Order of operations:
  (a) text edits to existing slides (run-level, formatting preserved);
  (b) append the 24 new slides from content_ch06_enrich;
  (c) ONE final arrange() placing every slide in its pedagogical position
      (old slide 21 — unverified hellotars/IOG examples — is dropped);
  (d) save, then verify (count, zip integrity).
"""
from pathlib import Path
import zipfile
import pptx_tools as P
from pptx.oxml.ns import qn
import content_ch06_enrich as C

DECK = Path("/Users/iahmad/Creator/Courses_and_conferences/LT/1258/1258a4-author-input/decks/1258-Ch06-Data-for-AI.pptx")


def clean_orphan_slides(prs):
    """The source deck carries orphan slide parts (slide51/52 + their notesSlides),
    leftover from its edit history and reachable only via viewProps.xml rels.
    Appended slides take partnames len(sldIdLst)+1 and collide with them on save,
    producing duplicate zip entries. Drop the rels so the orphans become unreachable."""
    real = {prs.part.rels[e.get(qn("r:id"))].target_part for e in prs.slides._sldIdLst}
    drops = []
    for pt in prs.part.package.iter_parts():
        for rId, rel in pt.rels.items():
            if rel.is_external:
                continue
            tp = rel.target_part
            if str(tp.partname).startswith("/ppt/slides/") and tp not in real:
                drops.append((pt, rId))
    for pt, rId in drops:
        if hasattr(pt, "drop_rel"):
            pt.drop_rel(rId)
        else:
            pt.rels.pop(rId)
    return len(drops)


def edit_bullet(slide, old_sub, new_sub):
    """Replace text inside a bullet, preserving the first run's formatting."""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            full = "".join(r.text for r in p.runs)
            if old_sub in full:
                newfull = full.replace(old_sub, new_sub)
                p.runs[0].text = newfull
                for r in p.runs[1:]:
                    r.text = ""
                return True
    return False


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def append_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = tf.text.rstrip() + "\n" + text


def text_edits(prs):
    s = list(prs.slides)

    # Slide 4 (idx 3): soften launch-marketing figure (Corrections #7)
    assert edit_bullet(s[3], "Over 250 built-in transformations",
                       "Hundreds of pre-built transformations")
    append_notes(s[3],
        "Correction (a3): the 'over 250' figure traces only to AWS's 2020 launch marketing; "
        "softened to 'hundreds.' DataBrew remains an offered AWS service as of 2026-08-01.\n"
        + C.DATABREW)

    # Slide 11 (idx 10): GDPR-led compliance note -> US federal obligations (Corrections #9)
    set_notes(s[10],
        "Speaker Notes:\n"
        "This slide highlights the key challenges government agencies face in managing their data effectively.\n"
        "Start by emphasizing the increasing volume, variety, and velocity of data in the public sector, creating unprecedented challenges for data management.\n"
        "Discuss the first challenge: volume. Explain how the massive influx of data from various sources, such as healthcare, transportation, and public safety, strains existing infrastructure. Highlight the difficulties in scaling infrastructure to accommodate this data growth, which can lead to processing delays and hinder timely analysis.\n"
        "Next, discuss the challenge of data quality. Explain how data silos impede collaboration and hinder cross-agency cooperation. Discuss how inconsistent data formats, missing information, and duplication can compromise the reliability of AI systems and the accuracy of insights.\n"
        "Finally, emphasize data privacy and security. For this US federal audience, lead with US obligations: the Privacy Act of 1974, FISMA, and CIPSEA (all enumerated in OMB M-25-05's applicability list); keep HIPAA for health-data examples and mention GDPR only as international context. Safeguarding sensitive government data requires robust protection measures and controls to prevent breaches and unauthorized access.\n"
        + C.M2505)

    # Slide 18 (idx 17): remove unsourced '30% error reduction' claim (Gaps)
    set_notes(s[17],
        "Speaker Notes:\n"
        "Highlight examples of AI reducing time spent on repetitive tasks, allowing staff to focus on strategic initiatives.\n"
        "An earlier version cited AI data-quality checks in a national benefits program with a specific "
        "error-reduction percentage — no source could be verified, so the figure was removed. "
        "Use only outcomes you can source.")

    # Slide 22 (idx 21): AI governance models — add federal anchors to notes
    set_notes(s[21],
        "Speaker Notes:\n"
        "This slide presents centralized/hybrid vs decentralized/federated AI governance models generically. "
        "Anchor it federally: NIST AI RMF's Govern function is the federal frame for AI oversight structure; "
        "OMB M-25-05 assigns data governance to agency Chief Data Officers and Data Governance Bodies; "
        "OMB M-25-21 adds minimum risk-management practices for high-impact AI. The next slides develop all three.\n"
        + C.NISTRMF + "\n" + C.M2505)

    # Slide 27 (idx 26): summary lists the wrong tools (Corrections #2)
    assert edit_bullet(s[26], "Tools explored: Tableau, Pandas, Scikit-learn, and AWS QuickSight for scalable and interactive analytics",
                       "Tools explored: AWS Glue DataBrew, SageMaker Data Wrangler, AWS Glue, Power BI, and Pandas")
    assert edit_bullet(s[26], "Ai solutions:", "AI solutions:")
    append_notes(s[26],
        "Correction (a3): the tools line previously listed Tableau/Scikit-learn/QuickSight, which this chapter "
        "does not teach; aligned with the chapter's actual tools (slides: DataBrew, Data Wrangler, AWS Glue, Power BI, Pandas).")

    # Slide 29 (idx 28): Ex 6.1 launcher — keep, add spine-rung P6 + six-dimension audit
    append_notes(s[28],
        "Prompting spine: this exercise is rung P6. Run the audit explicitly against the six dimensions "
        "(accuracy, completeness, consistency, timeliness, validity, uniqueness) — the three walkthrough "
        "slides immediately before this launcher script the flow.")


def arrange_order(n_old=29, n_new=24):
    """Full permutation: every old slide except dropped idx 20 (old slide 21),
    plus all new slides (idx 29..52), in final pedagogical order."""
    N = list(range(n_old, n_old + n_new))          # N[0..23] = new slides N1..N24
    order = [
        0, 1, 2, 10, 8,          # title, objectives, bottleneck, key challenges, prep journey
        N[0], N[1],              # collection methods, schema first
        9, 18,                   # data-quality assessment, intro to data quality
        N[2],                    # three frames
        N[3], N[4], N[5], N[6], N[7], N[8],   # six dimensions
        19,                      # relevance and accuracy (old slide 20)
        N[9], N[10], N[11],      # Ex 6.1 walkthrough x3
        28,                      # Ex 6.1 launcher (moved from deck end)
        11, 12, 13,              # What is GenAI, collection, cleansing
        N[12], N[13],            # structured outputs, validation habit
        3, 4, 5, 6,              # DataBrew x2, Data Wrangler x2
        N[14],                   # tool landscape [flex]
        7,                       # shift toward automated pipelines
        N[15], N[16],            # pipeline anatomy [flex], storage choices [flex]
        N[17], N[18], N[19], N[20],  # labeling, synthetic, ONS spectrum [flex], de-id [flex]
        14, 15, 16, 17,          # use cases 1-3, role of AI
        21,                      # AI governance for government
        N[21], N[22],            # federal data governance, federal AI oversight
        22, 23, 24, 25,          # data governance framework x4
        N[23],                   # datasheets/DCAT-US [flex]
        27,                      # objectives review (moved from deck end)
        26,                      # summary
    ]
    assert sorted(order) == sorted(set(range(n_old + n_new)) - {20}), \
        "order must be a permutation of all slides except dropped idx 20"
    return order


def build():
    prs = P.open_deck(DECK)
    n_dropped = clean_orphan_slides(prs)
    P.save(prs, DECK); prs = P.open_deck(DECK)   # save+reopen (partname hygiene)
    print(f"cleaned {n_dropped} orphan slide rels")
    text_edits(prs)
    for t, b, n in C.NEW_SLIDES:
        P.append_content(prs, t, b, n)
    P.arrange(prs, arrange_order())
    P.save(prs, DECK)
    return prs


if __name__ == "__main__":
    prs = build()
    n = len(list(prs.slides._sldIdLst))
    assert zipfile.ZipFile(DECK).testzip() is None, "zip integrity failed"
    names = zipfile.ZipFile(DECK).namelist()
    dupes = {x for x in names if names.count(x) > 1}
    assert not dupes, f"dup partnames: {dupes}"
    print(f"{DECK.name}: {n} slides, zip OK, no dup partnames")
    for i, t, w in P.inventory(DECK):
        print(f"{i:>3}  {w:>4}w  {t}")
