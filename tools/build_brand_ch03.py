"""
build_brand_ch03.py — apply LT standard branding/structure to
1258-Ch03-Desktop-GenAI.pptx (Chapter 3 'Desktop GenAI: Everyday AI for
Government Staff').

Idempotent: always rebuilds from decks/_bak6/1258-Ch03-Desktop-GenAI.pptx.

Changes (vs. 57-slide original):
  1. New Chapter Title slide (layout 'Chapter Title Government'):
     CENTER_TITLE (idx 0) = chapter name, SUBTITLE (idx 1) = 'Chapter 3'.
     The old slide 1 (a Content slide acting as the title, with the chapter
     objectives as its body) is folded into the new title slide's speaker
     notes and dropped.
  2. New Chapter Objectives slide (layout 'Chapter Objectives'), body idx 15.
  3. New Contents slide (layout 'Content with Header. Full Page'):
     section lines + the chapter's activities.
  4. Old slide 52 'Lab 3.1: Desktop GenAI Task Relay' rebuilt on layout
     'Exercise Reference Slide with Typing Hands' at the same position;
     body = Lab Manual reference + key lines incl. the registry time
     (30 min, registry/activities.yaml); speaker notes carried verbatim.
     The 'Lab 3.1 Walkthrough' / 'Lab 3.1 Debrief' slides are teaching
     content around the lab (not the exercise reference) and stay on
     'Content with Header. Full Page'. The two 'Activities in This Chapter'
     overview slides (registry-generated) are kept as content slides.
     Demo 3.1 / DO NOW 3.A-3.D have no dedicated slides in this deck —
     they exist only as lines on the overview slides + the Lab Manual.

Final slide count: 57 - 1 folded opener + 3 openers = 59 (in the 50-60 band;
no [flex] trims needed).

Run:  ../../.venv-courseware/bin/python tools/build_brand_ch03.py
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_tools import open_deck, arrange, save  # noqa: E402

DECKS = Path(__file__).resolve().parent.parent / "decks"
SRC = DECKS / "_bak6" / "1258-Ch03-Desktop-GenAI.pptx"
OUT = DECKS / "1258-Ch03-Desktop-GenAI.pptx"

CHAPTER_NAME = "Desktop GenAI: Everyday AI for Government Staff"
SUBTITLE = "Chapter 3"

OBJECTIVES = [
    "Recognize the four desktop assistants government staff use daily — and their government editions",
    "Choose the right assistant and edition by FedRAMP authorization status",
    "Apply GenAI to five everyday work patterns: research, summarize, draft, diagram, code help",
    "Apply the data rule and your agency's acceptable use policy (AUP) to every prompt",
    "Spot shadow-AI risk — and know when not to use a desktop assistant",
]

CONTENTS = [
    "The four assistants government staff use daily",
    "Government editions and the FedRAMP authorization scoreboard",
    "Five everyday work patterns, with worked examples",
    "Free, paid, and enterprise tiers — who trains on your prompts",
    "The data rule, your agency AUP, and shadow AI",
    "Activities: Demo 3.1 · DO NOW 3.A–3.D · Lab 3.1 (Desktop GenAI Task Relay)",
]

LAB_TITLE = "Lab 3.1: Desktop GenAI Task Relay"  # old slide 52
LAB_BODY = [
    "In your Lab Manual, refer to Lab 3.1: Desktop GenAI Task Relay",
    "Four timed tasks with a real assistant: summarize, diagram, troubleshoot, draft",
    "Public/synthetic inputs only; score each output with the 3-point rubric",
    "30 minutes, including debrief",
]


def _layout(prs, name):
    for lay in prs.slide_masters[0].slide_layouts:
        if lay.name == name:
            return lay
    raise KeyError(f"layout not found: {name!r}")


def _ph(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    raise KeyError(f"placeholder idx {idx} not on slide (layout {slide.slide_layout.name!r})")


def _fill_body(ph, lines):
    tf = ph.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line


def _slide_lines(slide):
    """Title text, body paragraphs (main non-title placeholder), notes."""
    title = slide.shapes.title.text_frame.text if slide.shapes.title else ""
    body = None
    for ph in slide.placeholders:
        if ph.placeholder_format.idx != 0 and ph.has_text_frame:
            if body is None or (ph.width or 0) * (ph.height or 0) > (body.width or 0) * (body.height or 0):
                body = ph
    bullets = [p.text for p in body.text_frame.paragraphs if p.text.strip()] if body is not None else []
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    return title, bullets, notes


def _clean_orphan_slides(prs):
    """Drop rels to slide parts that are not in sldIdLst (leftovers from
    earlier edit history collide with partnames python-pptx assigns to
    appended slides). Same pattern as a4_common.clean_orphan_slides."""
    from pptx.oxml.ns import qn

    real = {prs.part.rels[e.get(qn("r:id"))].target_part for e in prs.slides._sldIdLst}
    drops = []
    for pt in prs.part.package.iter_parts():
        for rId, rel in pt.rels.items():
            if rel.is_external:
                continue
            tp = rel.target_part
            if str(tp.partname).startswith("/ppt/slides/") and tp not in real:
                drops.append((pt, rId))
    for pt, rId in drops:
        if hasattr(pt, "drop_rel"):
            pt.drop_rel(rId)
        else:
            pt.rels.pop(rId)
    return len(drops)


def main():
    shutil.copy(SRC, OUT)
    prs = open_deck(OUT)
    slides = list(prs.slides)
    assert len(slides) == 57, f"expected 57 slides, got {len(slides)}"
    print(f"cleaned {_clean_orphan_slides(prs)} orphan slide rels")

    # Capture content to carry over before any rebuild.
    old1_title, old1_bullets, old1_notes = _slide_lines(slides[0])
    assert old1_title == CHAPTER_NAME, old1_title
    old52_title, _old52_bullets, old52_notes = _slide_lines(slides[51])
    assert old52_title == LAB_TITLE, old52_title

    # 1. Chapter Title (folds old slide 1: objectives body + notes -> notes)
    t = prs.slides.add_slide(_layout(prs, "Chapter Title Government"))
    _ph(t, 0).text_frame.text = CHAPTER_NAME
    _ph(t, 1).text_frame.text = SUBTITLE
    t.notes_slide.notes_text_frame.text = "\n".join(
        [old1_notes, "", "Folded from the original opener slide (now the Chapter Objectives slide):"]
        + old1_bullets
    ).strip()

    # 2. Chapter Objectives (idx 15 is the displayed body on this layout)
    o = prs.slides.add_slide(_layout(prs, "Chapter Objectives"))
    _fill_body(_ph(o, 15), OBJECTIVES)

    # 3. Contents
    c = prs.slides.add_slide(_layout(prs, "Content with Header. Full Page"))
    _ph(c, 0).text_frame.text = "Contents"
    _fill_body(_ph(c, 1), CONTENTS)

    # 4. Lab 3.1 launcher -> Exercise Reference Slide with Typing Hands
    lab = prs.slides.add_slide(_layout(prs, "Exercise Reference Slide with Typing Hands"))
    _ph(lab, 0).text_frame.text = LAB_TITLE
    _fill_body(_ph(lab, 1), LAB_BODY)
    lab.notes_slide.notes_text_frame.text = old52_notes

    # Order: title, objectives, contents, orig 2..51, lab, orig 53..57
    # (0-based orig indices 0..56; new slides appended at 57,58,59,60)
    order0 = [57, 58, 59] + list(range(1, 51)) + [60] + list(range(52, 57))
    arrange(prs, order0)
    save(prs, OUT)

    # quick echo
    prs2 = open_deck(OUT)
    print(f"saved {OUT} ({len(list(prs2.slides))} slides)")
    for i, s in enumerate(prs2.slides):
        title = s.shapes.title.text_frame.text.replace("\n", " ")[:55] if s.shapes.title else "(no title)"
        print(f"{i+1:>2} [{s.slide_layout.name}] {title}")


if __name__ == "__main__":
    main()
