"""
notes_common.py — shared helpers for the delivery-notes rewrite (build_notes_*.py).

Each build_notes_<tag>.py is idempotent: it restores its deck from decks/_bak7/
(the pre-rewrite backup), rebuilds every slide's notes in delivery format
(SAY/ASK hook, KEY POINT, time cue, TRANSITION), and re-appends the preserved
markers extracted verbatim from the backup notes:

  - [flex] prefix semantics: slides whose backup notes start with "[flex]" get
    new notes starting with "[flex] (skip if behind — core message: ...)".
    The OLD flex prose is replaced by this standardized prefix, but Source
    citations and AUTHOR FLAG statements inside flex lines are still extracted.
  - "Source: ... (verified 2026-08-01)" / "Sources: ..." citation lines:
    the substring from "Source(s):" to end of line, kept exactly.
  - "Prompting spine: rung Pn." lines (any line mentioning rung P<n>): whole line.
  - Author/IG flag lines: [TEACH] registry/launcher lines, AUTHOR FLAG / AUTHOR/IG
    statements, "generated from the activity registry", verify-before-class and
    recheck cautions, and author correction/do-not-claim cautions
    (Correction/CORRECTED/Sourcing note/Terminology update/Currency/unsourced/
    unverified/fabricated/do not quote). Kept exactly (whole line).

verify() re-checks a finished deck: slide count unchanged, zip integrity,
slide body text unchanged vs. backup, the 4 delivery elements present in every
note, and every preserved marker present verbatim.
"""
from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation

DECKS = Path(__file__).resolve().parent.parent / "decks"
BAK = DECKS / "_bak7"

# Substrings (matched case-insensitively) that mark a notes line as an
# author/IG flag to preserve verbatim.
FLAG_SUBSTRS = (
    "[teach]",
    "generated from the activity registry",
    "re-check", "recheck", "re-checked",
    "verify before", "verify-before",
    "do not quote",
    "correction", "corrected", "corrections",
    "sourcing note", "terminology update", "currency",
    "unsourced", "unverified", "unverifiable", "not verifiable",
    "fabricated", "no supporting source", "could not be verified",
)
# Flags preserved from the keyword to end of line (so leading flex prose is dropped).
AUTHOR_KWS = ("AUTHOR FLAG:", "AUTHOR/IG:")

SRC_RE = re.compile(r"\bSources?:\s")
RUNG_RE = re.compile(r"rung\s+P\d+", re.IGNORECASE)
TIME_RE = re.compile(r"^\([^\n]*\d+[\s-]*min[^\n]*\)$", re.MULTILINE)


def extract_preserved(notes: str):
    """Extract (is_flex, kept_marker_lines) from a slide's ORIGINAL notes text."""
    kept = []
    for raw in notes.split("\n"):
        line = raw.strip()
        if not line:
            continue
        is_flexline = line.startswith("[flex]")
        captured = False
        for kw in AUTHOR_KWS:
            if kw in line:
                kept.append(line[line.index(kw):])
                captured = True
                break
        if not captured and not is_flexline and any(p in line.lower() for p in FLAG_SUBSTRS):
            kept.append(line)
            captured = True
        m = SRC_RE.search(line)
        if m and "verified" in line:
            src = line[m.start():]
            if not captured or not kept or src not in kept[-1]:
                kept.append(src)
        if RUNG_RE.search(line) and not captured:
            kept.append(line)
    seen, out = set(), []
    for k in kept:
        if k not in seen:
            seen.add(k)
            out.append(k)
    return notes.strip().startswith("[flex]"), out


def compose(delivery: str, flex_core: str | None, kept: list[str]) -> str:
    blocks = []
    if flex_core is not None:
        blocks.append(f"[flex] (skip if behind — core message: {flex_core})")
    blocks.append(delivery.strip())
    text = "\n".join(blocks)
    if kept:
        text += "\n\n" + "\n".join(kept)
    return text


def apply_notes(tag: str, expected_count: int, notes: dict, flex_core: dict):
    """Restore decks/1258-<tag>.pptx from _bak7 and rewrite all notes. Returns path."""
    deck = DECKS / f"1258-{tag}.pptx"
    bak = BAK / f"1258-{tag}.pptx"
    shutil.copy(bak, deck)  # idempotent rebuild: always start from the backup
    prs = Presentation(deck)
    slides = list(prs.slides)
    assert len(slides) == expected_count, f"{tag}: {len(slides)} slides != expected {expected_count}"
    assert set(notes) == set(range(1, len(slides) + 1)), (
        f"{tag}: NOTES keys {sorted(set(range(1, len(slides)+1)) - set(notes))} missing, "
        f"extras {sorted(set(notes) - set(range(1, len(slides)+1)))}")
    flex_detected = []
    for i, slide in enumerate(slides, 1):
        old = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        is_flex, kept = extract_preserved(old)
        if is_flex:
            flex_detected.append(i)
            assert i in flex_core, f"{tag} s{i}: flex slide missing FLEX_CORE entry"
        else:
            assert i not in flex_core, f"{tag} s{i}: FLEX_CORE entry but backup notes are not [flex]"
        slide.notes_slide.notes_text_frame.text = compose(
            notes[i], flex_core.get(i) if is_flex else None, kept)
    prs.save(deck)
    print(f"{tag}: {len(slides)} slide notes rewritten; flex slides: {flex_detected}")
    return deck


def verify(tag: str, expected_count: int) -> int:
    """Assert count/zip/slide-text unchanged, 4 elements per note, markers preserved.
    Returns the number of preserved markers verified."""
    deck = DECKS / f"1258-{tag}.pptx"
    bak = BAK / f"1258-{tag}.pptx"
    assert zipfile.ZipFile(deck).testzip() is None, f"{tag}: zipfile.testzip() failed"
    prs, bprs = Presentation(deck), Presentation(bak)
    slides, bslides = list(prs.slides), list(bprs.slides)
    assert len(slides) == expected_count == len(bslides), (
        f"{tag}: slide count {len(slides)} (backup {len(bslides)}) != {expected_count}")
    n_kept = 0
    for i, (s, bs) in enumerate(zip(slides, bslides), 1):
        st = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        bt = [sh.text_frame.text for sh in bs.shapes if sh.has_text_frame]
        assert st == bt, f"{tag} s{i}: slide body text changed"
        note = s.notes_slide.notes_text_frame.text
        bold = bs.notes_slide.notes_text_frame.text if bs.has_notes_slide else ""
        assert re.search(r"^(SAY|ASK):", note, re.MULTILINE), f"{tag} s{i}: missing SAY/ASK hook"
        assert re.search(r"^KEY POINT:", note, re.MULTILINE), f"{tag} s{i}: missing KEY POINT"
        assert TIME_RE.search(note), f"{tag} s{i}: missing (…min…) time cue"
        assert re.search(r"^TRANSITION:", note, re.MULTILINE), f"{tag} s{i}: missing TRANSITION"
        is_flex, kept = extract_preserved(bold)
        if is_flex:
            assert note.startswith("[flex] (skip if behind — core message:"), (
                f"{tag} s{i}: flex notes must start with '[flex] (skip if behind — core message: …)'")
        for k in kept:
            assert k in note, f"{tag} s{i}: preserved marker missing: {k[:70]}"
            n_kept += 1
    print(f"{tag}: OK — {len(slides)} slides, zip clean, slide text unchanged, "
          f"4 elements in every note, {n_kept} preserved markers verified")
    return n_kept
